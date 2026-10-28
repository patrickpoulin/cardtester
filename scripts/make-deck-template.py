"""Writes the deck templates in public/templates/: each game's original deck plus a sheet explaining the columns.

deck-template.xlsx is Not Snap's, rascally-raccoons-deck-template.xlsx is Rascally Raccoons'. The rows come from
the rules files in src/games/ (tests/deck.test.ts and tests/raccoonDeck.test.ts check they stay in sync). Needs openpyxl.
Run from the repository root: npx tsx scripts/deck-template-data.ts | python3 scripts/make-deck-template.py
"""

import json, csv, io, sys, openpyxl
from openpyxl.styles import Font


def write(path, data, int_cols, widths, columns, effect_heading, notes):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'Deck'
    for i, row in enumerate(csv.reader(io.StringIO(data['csv']))):
        if not row:
            continue
        out = []
        for j, v in enumerate(row):
            if i > 0 and j in int_cols and v != '':
                v = int(v)
            out.append(v if v != '' else None)
        ws.append(out)
    for c in ws[1]:
        c.font = Font(bold=True)
    for col, w in zip('ABCDEFG', widths):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = 'A2'
    h = wb.create_sheet('How to fill this in')
    h.append(['Column', 'What it means'])
    for c in h[1]:
        c.font = Font(bold=True)
    for r in columns:
        h.append(list(r))
    h.append([])
    h.append(effect_heading)
    for c in h[h.max_row]:
        c.font = Font(bold=True)
    for e in data['effects']:
        h.append([e['label'], e['description']])
    h.append([])
    for n in notes:
        h.append(['', n])
    h.append(['', 'Card Tester reads only the first sheet (Deck).'])
    h.column_dimensions['A'].width = 22
    h.column_dimensions['B'].width = 90
    wb.save(path)


d = json.load(sys.stdin)
write('public/templates/deck-template.xlsx', d['notSnap'], (1, 2, 3, 5), (16, 11, 7, 7, 15, 9, 52), [
    ('Name', "The card's name. Required, and each name only once."),
    ('Per Player', "Copies of the card in each player's deck. Empty means 1."),
    ('Cost', 'Energy needed to play it (a whole number). Empty counts as 0.'),
    ('Power', 'Its power at a location (a whole number). Empty counts as 0.'),
    ('Effect', 'Optional ability: one of the effects below. Leave empty for no ability.'),
    ('Amount', 'The number the Effect uses. Required when Effect is set.'),
    ('Ability', 'The card text, for people reading the deck. Not used by the simulator.'),
], ['Effect', 'What it does'], [])
write('public/templates/rascally-raccoons-deck-template.xlsx', d['raccoons'], (2,), (24, 70, 8, 24), [
    ('Card Name', "The card's name. Required, and each name only once."),
    ('Text', 'The card text, for people reading the deck. Not used by the simulator.'),
    ('Copies', 'How many of the card are in the deck. Empty means 1.'),
    ('Effect', 'Which card below it plays like. Can be left empty when the Card Name is one of them.'),
], ['Effect', 'What it does'], [
    'The deck needs Percy the Dog, at least one Rascally Raccoon, and at least 11 other cards (35 for 6 players).',
])
